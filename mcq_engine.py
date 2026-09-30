import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel
from typing import List, Literal
import json
import pymupdf4llm
import pypandoc
import os
import shutil
import re
import base64
import mimetypes
from docx import Document
from docx.shared import Mm, Inches
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# 1. STRUCTURED OUTPUT SCHEMA
# ============================================================
# Gemini's structured-output feature uses this schema to keep the
# response machine-readable. This removes most JSON / LaTeX escape
# failures before they reach json.loads().

class MCQItem(BaseModel):
    sl_no: int
    question_title: str
    A: str
    B: str
    C: str
    D: str
    solution_body: str
    correct_option: Literal["A", "B", "C", "D"]
    subject: str
    chapter: str
    topics: str
    question_category: str
    difficulty_level: Literal["Easy", "Medium", "Hard"]


class MCQResponse(BaseModel):
    questions: List[MCQItem]


# Keep individual model responses reasonably sized. Large single JSON
# responses are more likely to be truncated or malformed.
GENERATION_BATCH_SIZE = 15
MODEL_RETRY_ATTEMPTS = 2


# ============================================================
# 2. SOURCE EXTRACTION
# ============================================================

def extract_text_from_pdf(pdf_file):
    """Extract PDF text as Markdown and preserve images in ./media."""
    if os.path.exists("media"):
        shutil.rmtree("media")
    if os.path.exists("word"):
        shutil.rmtree("word")

    os.makedirs("media", exist_ok=True)

    with open("temp.pdf", "wb") as f:
        f.write(pdf_file.read())

    try:
        text = pymupdf4llm.to_markdown(
            "temp.pdf",
            write_images=True,
            image_path="media",
        )
    finally:
        if os.path.exists("temp.pdf"):
            os.remove("temp.pdf")

    return normalize_extracted_source(text)


def extract_text_from_docx(docx_file):
    """Extract DOCX text as Markdown and preserve embedded media."""
    if os.path.exists("media"):
        shutil.rmtree("media")
    if os.path.exists("word"):
        shutil.rmtree("word")

    with open("temp.docx", "wb") as f:
        f.write(docx_file.read())

    try:
        text = pypandoc.convert_file(
            "temp.docx",
            "markdown",
            extra_args=["--extract-media=."],
        )
    finally:
        if os.path.exists("temp.docx"):
            os.remove("temp.docx")

    return normalize_extracted_source(text)


def normalize_extracted_source(text):
    """
    Clean extraction noise without changing the actual educational content.

    Important:
    - Markdown image paths are preserved.
    - Pandoc width/height attribute blocks are removed because they often
      confuse the model and are not needed for later DOCX rendering.
    - Excess blank lines and non-breaking spaces are normalized.
    """
    if text is None:
        return ""

    text = str(text).replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ")

    # Preserve the image Markdown itself but remove a following Pandoc
    # dimensions block such as {width="2in" height="1in"}.
    text = re.sub(
        r'(!\[[^\]]*\]\([^)]+\))\s*\{[^{}]*(?:width|height)[^{}]*\}',
        r'\1',
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    # Remove orphaned image dimension blocks left by conversion.
    text = re.sub(
        r'\{[^{}]*(?:width|height)[^{}]*\}',
        '',
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    # Avoid huge blank areas in the model prompt.
    text = re.sub(r'\n{4,}', '\n\n\n', text)

    return text.strip()


# ============================================================
# 3. JSON / STRUCTURED RESPONSE SAFETY HELPERS
# ============================================================

def strip_markdown_fences(text):
    """Remove accidental ```json / ``` wrappers."""
    text = str(text or "").strip()

    if text.startswith("```json"):
        text = text[7:]
    elif text.startswith("```"):
        text = text[3:]

    if text.rstrip().endswith("```"):
        text = text.rstrip()[:-3]

    return text.strip()


def isolate_json_candidate(text):
    """
    If a model accidentally adds commentary, keep the most likely JSON payload.
    Structured output normally makes this unnecessary, but it is a safe fallback.
    """
    text = strip_markdown_fences(text)

    first_array = text.find("[")
    first_object = text.find("{")

    starts = [x for x in (first_array, first_object) if x >= 0]
    if not starts:
        return text

    start = min(starts)
    opening = text[start]
    closing = "]" if opening == "[" else "}"
    end = text.rfind(closing)

    if end > start:
        return text[start:end + 1].strip()

    return text[start:].strip()



def strip_raw_pandoc_dimension_blocks(text):
    """
    Remove raw Pandoc {width="..." height="..."} artifacts before JSON
    parsing. If Gemini copies these attributes into a JSON string without
    escaping the inner quotes, the entire JSON becomes invalid. Image paths are
    unaffected; sizing is handled later by process_html_images().
    """
    return re.sub(
        r'\{[^{}\n]*(?:width|height)\s*=\s*"[^"]*"[^{}\n]*\}',
        '',
        text,
        flags=re.IGNORECASE,
    )


def escape_literal_control_chars_inside_json_strings(text):
    """
    Convert literal newlines/tabs inside quoted JSON strings into escaped JSON
    sequences. This repairs another common failure mode in long math outputs.
    """
    output = []
    in_string = False
    escaped = False

    for ch in text:
        if in_string:
            if escaped:
                output.append(ch)
                escaped = False
                continue

            if ch == "\\":
                output.append(ch)
                escaped = True
                continue

            if ch == '"':
                output.append(ch)
                in_string = False
                continue

            if ch == "\n":
                output.append("\\n")
                continue
            if ch == "\r":
                output.append("\\r")
                continue
            if ch == "\t":
                output.append("\\t")
                continue

            output.append(ch)
        else:
            output.append(ch)
            if ch == '"':
                in_string = True
                escaped = False

    return "".join(output)


def repair_latex_json_escapes(text):
    r"""
    Protect common LaTeX commands that JSON would otherwise mistake for valid
    one-character escapes. For example, raw JSON "\frac" can be interpreted as
    JSON form-feed (\f) + "rac", and "\theta" as tab (\t) + "heta".

    Correctly escaped JSON LaTeX ("\\frac") is left untouched.
    """
    latex_commands = (
        "frac|dfrac|tfrac|sqrt|angle|triangle|bigtriangleup|left|right|times|"
        "text|therefore|because|cdot|circ|degree|alpha|beta|gamma|delta|theta|"
        "lambda|mu|pi|rho|sigma|phi|omega|sin|cos|tan|cot|sec|csc|log|ln|"
        "sum|prod|int|lim|infty|overline|underline|underbrace|overbrace|vec|"
        "hat|bar|mathbf|mathrm|operatorname|begin|end|pm|mp|leq|geq|neq|"
        "approx|equiv|in|notin|subset|supset|cup|cap|perp|parallel"
    )

    return re.sub(
        rf'(?<!\\)\\(?=(?:{latex_commands})\b)',
        r'\\\\',
        text,
        flags=re.IGNORECASE,
    )


def repair_invalid_json_escapes(text):
    r"""
    Escape remaining backslashes that are illegal in JSON strings.

    Valid JSON escapes are:
        \"  \\  \/  \b  \f  \n  \r  \t  \uXXXX
    """
    return re.sub(
        r'(?<!\\)\\(?!["\\/bfnrt]|u[0-9a-fA-F]{4})',
        r'\\\\',
        text,
    )


def robust_json_loads(raw_output):
    """
    Multi-stage fallback JSON parser.

    Primary flow uses Gemini response.parsed, so this runs only when needed.
    """
    candidate = isolate_json_candidate(raw_output)
    errors = []

    candidate = strip_raw_pandoc_dimension_blocks(candidate)
    controls_repaired = escape_literal_control_chars_inside_json_strings(candidate)
    latex_repaired = repair_latex_json_escapes(controls_repaired)
    repaired = repair_invalid_json_escapes(latex_repaired)

    # Parse the protected version first. This avoids silent corruption such as
    # JSON treating LaTeX \frac as the valid JSON escape \f + "rac".
    attempts = [repaired, latex_repaired, controls_repaired]

    # Avoid trying identical strings repeatedly.
    unique_attempts = []
    for attempt in attempts:
        if attempt not in unique_attempts:
            unique_attempts.append(attempt)

    for attempt in unique_attempts:
        try:
            return json.loads(attempt)
        except json.JSONDecodeError as e:
            errors.append(str(e))

    # Last-resort relaxed parser for control characters.
    try:
        return json.loads(repaired, strict=False)
    except json.JSONDecodeError as e:
        errors.append(str(e))

    raise ValueError(
        "JSON could not be parsed after automatic repair. "
        f"Last parser error: {errors[-1] if errors else 'Unknown JSON error'}"
    )


def pydantic_to_dict(item):
    """Support both Pydantic v2 and v1 style model serialization."""
    if isinstance(item, dict):
        return item
    if hasattr(item, "model_dump"):
        return item.model_dump()
    if hasattr(item, "dict"):
        return item.dict()
    raise TypeError(f"Unsupported structured item type: {type(item)}")


def payload_to_records(payload):
    """Convert Gemini parsed/raw payload into a plain list[dict]."""
    if isinstance(payload, MCQResponse):
        items = payload.questions
    elif isinstance(payload, dict):
        if "questions" in payload:
            items = payload["questions"]
        else:
            # Graceful fallback for a single object.
            items = [payload]
    elif isinstance(payload, list):
        items = payload
    elif hasattr(payload, "questions"):
        items = payload.questions
    else:
        raise TypeError(f"Unexpected model payload type: {type(payload)}")

    return [pydantic_to_dict(item) for item in items]


def response_to_records(response):
    """Prefer response.parsed. Fall back to repaired response.text if needed."""
    parsed = getattr(response, "parsed", None)

    if parsed is not None:
        try:
            return payload_to_records(parsed)
        except Exception:
            # If SDK parsing produced an unexpected object, try response.text.
            pass

    raw_text = getattr(response, "text", None)
    if not raw_text:
        raise ValueError("Gemini returned no parseable structured data or text.")

    return payload_to_records(robust_json_loads(raw_text))


def get_finish_reason(response):
    """Best-effort extraction of Gemini finish reason for clearer errors."""
    try:
        if response.candidates:
            reason = response.candidates[0].finish_reason
            return str(reason) if reason is not None else ""
    except Exception:
        pass
    return ""


# ============================================================
# 4. MCQ NORMALIZATION / VALIDATION
# ============================================================

def split_topics(user_topics):
    if not user_topics:
        return []
    return [topic.strip() for topic in str(user_topics).split(",") if topic.strip()]


def clean_math_backslashes(obj):
    """
    Recursively reduce accidentally doubled LaTeX backslashes to one.
    JSON decoding itself already removes JSON-level escaping; this only cleans
    model-produced textual double escaping if it remains.
    """
    if isinstance(obj, str):
        return obj.replace("\\\\", "\\")
    if isinstance(obj, list):
        return [clean_math_backslashes(item) for item in obj]
    if isinstance(obj, dict):
        return {k: clean_math_backslashes(v) for k, v in obj.items()}
    return obj


def format_option_labels(parsed_data):
    """
    Ensure options use 'A. ...' / 'B. ...' format and correct_option is a
    single A/B/C/D letter. Also supports fallback model outputs where the
    correct answer is returned as full option text.
    """
    if not isinstance(parsed_data, list):
        return parsed_data

    for mcq in parsed_data:
        if not isinstance(mcq, dict):
            continue

        # Normalize correct_option first.
        if "correct_option" in mcq and isinstance(mcq["correct_option"], str):
            co_val = mcq["correct_option"].strip()

            match = re.match(
                r'^(?:Option\s*)?([A-D])(?:[\.\)\-:\s]|$)',
                co_val,
                re.IGNORECASE,
            )

            if match:
                mcq["correct_option"] = match.group(1).upper()
            else:
                # If the model returned the answer text instead of the letter,
                # match it against the four options.
                for opt in ["A", "B", "C", "D"]:
                    if opt in mcq and isinstance(mcq[opt], str):
                        raw_opt_val = re.sub(
                            rf'^{opt}[\.\)\-]\s*',
                            '',
                            mcq[opt],
                            flags=re.IGNORECASE,
                        ).strip()

                        if raw_opt_val == co_val or mcq[opt].strip() == co_val:
                            mcq["correct_option"] = opt
                            break

        # Normalize option labels.
        for opt in ["A", "B", "C", "D"]:
            value = mcq.get(opt, "")
            if not isinstance(value, str):
                value = str(value)

            clean_val = re.sub(
                rf'^{opt}[\.\)\-]\s*',
                '',
                value,
                flags=re.IGNORECASE,
            ).strip()

            mcq[opt] = f"{opt}. {clean_val}"

    return parsed_data


def normalize_topic_values(records, allowed_topics):
    """
    Normalize case/whitespace when the model returns an allowed topic with a
    small formatting difference. This never invents a new topic.
    """
    if not allowed_topics:
        return records

    lookup = {topic.casefold(): topic for topic in allowed_topics}

    for mcq in records:
        value = str(mcq.get("topics", "")).strip()
        normalized = lookup.get(value.casefold())
        if normalized is not None:
            mcq["topics"] = normalized

    return records


def normalize_mcq_records(records, allowed_topics=None):
    """Final common normalization applied to both tabs."""
    records = clean_math_backslashes(records)
    records = format_option_labels(records)
    records = normalize_topic_values(records, allowed_topics or [])

    # Guarantee sequential numbering after batching / retries.
    for index, mcq in enumerate(records, start=1):
        mcq["sl_no"] = index

    return records


def validate_required_fields(records):
    required = [
        "sl_no",
        "question_title",
        "A",
        "B",
        "C",
        "D",
        "solution_body",
        "correct_option",
        "subject",
        "chapter",
        "topics",
        "question_category",
        "difficulty_level",
    ]

    for idx, mcq in enumerate(records, start=1):
        if not isinstance(mcq, dict):
            raise ValueError(f"Question {idx} is not a valid MCQ object.")

        missing = [key for key in required if key not in mcq]
        if missing:
            raise ValueError(
                f"Question {idx} is missing required fields: {', '.join(missing)}"
            )

        if mcq.get("correct_option") not in ["A", "B", "C", "D"]:
            raise ValueError(
                f"Question {idx} has invalid correct_option: "
                f"{mcq.get('correct_option')!r}"
            )

        if mcq.get("difficulty_level") not in ["Easy", "Medium", "Hard"]:
            raise ValueError(
                f"Question {idx} has invalid difficulty_level: "
                f"{mcq.get('difficulty_level')!r}"
            )


def validate_allowed_topics(records, allowed_topics):
    if not allowed_topics:
        return

    allowed = set(allowed_topics)
    invalid = []

    for mcq in records:
        topic = mcq.get("topics", "")
        if topic not in allowed:
            invalid.append((mcq.get("sl_no", "?"), topic))

    if invalid:
        preview = ", ".join(
            f"Q{sl}: {topic!r}" for sl, topic in invalid[:5]
        )
        raise ValueError(
            "Model used topic names outside the allowed list. "
            f"Examples: {preview}"
        )


def validate_generation_counts(records, num_easy, num_medium, num_hard):
    expected_total = num_easy + num_medium + num_hard

    if len(records) != expected_total:
        raise ValueError(
            f"Expected {expected_total} questions but model returned {len(records)}."
        )

    actual = {"Easy": 0, "Medium": 0, "Hard": 0}
    for mcq in records:
        difficulty = mcq.get("difficulty_level")
        if difficulty in actual:
            actual[difficulty] += 1

    expected = {
        "Easy": num_easy,
        "Medium": num_medium,
        "Hard": num_hard,
    }

    if actual != expected:
        raise ValueError(
            "Difficulty distribution mismatch. "
            f"Expected {expected}, received {actual}."
        )


def validate_mcqs(records, allowed_topics=None):
    validate_required_fields(records)
    validate_allowed_topics(records, allowed_topics or [])


# ============================================================
# 5. GEMINI REQUEST ENGINE
# ============================================================

def make_generation_batches(num_easy, num_medium, num_hard, batch_size=GENERATION_BATCH_SIZE):
    """Split a large generation job into smaller exact-difficulty batches."""
    remaining = {
        "Easy": int(num_easy),
        "Medium": int(num_medium),
        "Hard": int(num_hard),
    }

    batches = []

    while sum(remaining.values()) > 0:
        batch = {"Easy": 0, "Medium": 0, "Hard": 0}
        slots = batch_size

        for level in ["Easy", "Medium", "Hard"]:
            if slots <= 0:
                break
            take = min(remaining[level], slots)
            batch[level] = take
            remaining[level] -= take
            slots -= take

        batches.append(batch)

    return batches


def call_structured_mcq_model(
    client,
    model,
    prompt,
    allowed_topics=None,
    temperature=0.1,
    expected_counts=None,
    max_attempts=MODEL_RETRY_ATTEMPTS,
):
    """
    Call Gemini with a Pydantic response schema, normalize the result, validate
    it, and automatically retry once with the validation error as correction
    context if needed.
    """
    last_error = None
    retry_note = ""

    for attempt in range(1, max_attempts + 1):
        attempt_prompt = prompt

        if retry_note:
            attempt_prompt += (
                "\n\nIMPORTANT RETRY CORRECTION:\n"
                "Your previous response failed application validation for this reason:\n"
                f"{retry_note}\n"
                "Return the complete result again and correct that issue."
            )

        try:
            response = client.models.generate_content(
                model=model,
                contents=attempt_prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=MCQResponse,
                    temperature=temperature,
                ),
            )

            records = response_to_records(response)
            records = normalize_mcq_records(records, allowed_topics or [])
            validate_mcqs(records, allowed_topics or [])

            if expected_counts is not None:
                validate_generation_counts(
                    records,
                    expected_counts.get("Easy", 0),
                    expected_counts.get("Medium", 0),
                    expected_counts.get("Hard", 0),
                )

            return records

        except Exception as e:
            last_error = e
            retry_note = str(e)

            # If the response was truncated, expose this useful detail in the
            # retry note. Smaller generation batches already reduce this risk.
            try:
                finish_reason = get_finish_reason(response)
                if finish_reason:
                    retry_note += f" | Gemini finish reason: {finish_reason}"
            except Exception:
                pass

    raise ValueError(
        "The model response could not be validated after automatic retry. "
        f"Last error: {last_error}"
    )


# ============================================================
# 6. MCQ GENERATION (TAB 1)
# ============================================================

def build_generation_prompt(
    context_text,
    allowed_topics,
    custom_instructions,
    num_easy,
    num_medium,
    num_hard,
):
    total_questions = num_easy + num_medium + num_hard

    if context_text.strip():
        source_instruction = (
            "Use the provided Context Text as the source material together with "
            "the user's instructions."
        )
        context_block = f"Context Text:\n{context_text}"
    else:
        source_instruction = (
            "No source text was provided. Use expert educational knowledge and "
            "the user's instructions."
        )
        context_block = "No context text provided."

    topics_text = ", ".join(allowed_topics)

    return f"""
Act as an expert educator and create exactly {total_questions} board-standard multiple-choice questions.

SOURCE RULE:
{source_instruction}

SPECIAL USER GENERATION INSTRUCTIONS (CRITICAL):
{custom_instructions if custom_instructions else "None provided. Follow standard board-level question generation."}

MATHEMATICS / SCIENCE FORMATTING:
- Preserve or write mathematical equations, physics expressions and chemistry formulas using standard LaTeX.
- Use $...$ for inline math and $$...$$ for display math where appropriate.
- Use normal LaTeX commands such as \\frac, \\sqrt, \\angle, \\triangle, \\left, \\right, \\times and \\text.
- Do not manually perform JSON escaping. The application's structured-output layer handles JSON serialization.

IMAGE RULES:
- If the Context Text contains Markdown image links such as ![image](media/img.png) or ![alt](word/media/image1.jpeg), preserve the image link exactly in the relevant field.
- Never invent, translate, rename or delete an existing image path.

DATA TABLE RULES:
- If the source contains a Markdown data table, convert it to a basic HTML <table>...</table> string in the relevant field.
- Do not return a Markdown table inside an MCQ field.
- Keep each HTML table continuous; do not intentionally insert line breaks inside the table markup.

DIFFICULTY DISTRIBUTION (EXACT):
- Easy: {num_easy}
- Medium: {num_medium}
- Hard: {num_hard}
- Use only the exact labels Easy, Medium and Hard.

TOPIC MAPPING (EXACT):
Allowed topics: {topics_text}
For each question, the topics value must be exactly ONE item from this allowed list.
Do not invent another topic and do not add sequence identifiers such as 1.1, 2.1, etc.

SOLUTION BODY FORMAT:
- Line 1: exact text of the correct option, without A./B./C./D. or another option label.
- Line 2: blank line.
- Line 3 onward: explanation beginning exactly with "ব্যাখ্যা:".
- The explanation should be clear and sufficiently detailed.

FIELD EXPECTATIONS:
- subject: appropriate subject name.
- chapter: appropriate chapter name.
- question_category: appropriate category such as Board, Practice, Model Test, etc.
- correct_option: only A, B, C or D
- The application's response schema defines the exact output fields. Return only the structured result; no commentary.

{context_block}
""".strip()


def generate_mcqs(
    context_text,
    user_topics,
    custom_instructions,
    num_easy,
    num_medium,
    num_hard,
    selected_model,
    api_key,
):
    client = genai.Client(api_key=api_key)
    allowed_topics = split_topics(user_topics)

    if not allowed_topics:
        raise ValueError("At least one allowed topic is required for generation.")

    batches = make_generation_batches(num_easy, num_medium, num_hard)
    all_records = []

    for batch_index, batch in enumerate(batches, start=1):
        prompt = build_generation_prompt(
            context_text=context_text,
            allowed_topics=allowed_topics,
            custom_instructions=custom_instructions,
            num_easy=batch["Easy"],
            num_medium=batch["Medium"],
            num_hard=batch["Hard"],
        )

        # Tell later batches not to intentionally duplicate earlier items.
        if batch_index > 1:
            prompt += (
                "\n\nThis is a later batch of the same generation job. "
                "Create fresh questions and avoid obvious duplicates of earlier batches."
            )

        batch_records = call_structured_mcq_model(
            client=client,
            model=selected_model,
            prompt=prompt,
            allowed_topics=allowed_topics,
            temperature=0.2,
            expected_counts=batch,
        )

        all_records.extend(batch_records)

    # Re-number after combining batches and re-check the final totals.
    all_records = normalize_mcq_records(all_records, allowed_topics)
    validate_mcqs(all_records, allowed_topics)
    validate_generation_counts(
        all_records,
        int(num_easy),
        int(num_medium),
        int(num_hard),
    )

    return all_records


# ============================================================
# 7. EXISTING MCQ PARSER / FORMATTER (TAB 2)
# ============================================================

def build_existing_mcq_prompt(
    raw_mcq_text,
    allowed_topics,
    special_instructions,
):
    if allowed_topics:
        topic_rule = (
            "Allowed topics: "
            + ", ".join(allowed_topics)
            + "\nFor every MCQ, topics must be exactly ONE item from this list. "
              "Do not invent another topic or add numbering such as 1.1 or 2.1."
        )
    else:
        topic_rule = (
            "No target topic list was supplied. Infer one concise, academically "
            "appropriate topic name for each MCQ."
        )

    return f"""
You are an expert educational content parser and converter.
Convert ALL MCQs found in the supplied raw content into the application's structured MCQ format.
Do not intentionally omit any question.

PARSING RULES:
- Preserve the original question meaning, options and supplied correct answer.
- Number the final MCQs sequentially starting from 1.
- Extract Question Title, Option A, Option B, Option C, Option D, Correct Option and Solution Body.
- If a solution/explanation is missing, create a brief accurate explanation UNLESS the special user instructions explicitly say not to generate an explanation.
- Infer appropriate subject, chapter, question_category and difficulty_level (Easy, Medium or Hard) when missing.

MATHEMATICS / SCIENCE FORMATTING:
- Preserve mathematical equations, chemical formulas and physics expressions using standard LaTeX.
- Use $...$ for inline math and $$...$$ for display math where appropriate.
- Use normal LaTeX commands such as \\frac, \\sqrt, \\angle, \\triangle, \\left, \\right, \\times and \\text.
- Do not manually perform JSON escaping. The application's structured-output layer handles JSON serialization.

IMAGE RULES:
- Preserve every existing Markdown image link exactly, including paths such as ![image](media/img.png) or ![alt](word/media/image1.jpeg).
- Put an image link in the field where the image belongs, usually question_title or solution_body.
- Never invent, translate, rename or delete an existing image path.

DATA TABLE RULES:
- Convert source Markdown tables to basic HTML <table>...</table> strings in the relevant field.
- Do not return Markdown tables inside MCQ fields.
- Keep each HTML table continuous without intentional line breaks inside its markup.

TOPIC RULE:
{topic_rule}

SOLUTION BODY FORMAT:
Unless the special user instructions explicitly request no explanation:
- Line 1: exact text of the correct option, without A./B./C./D. or another option label.
- Line 2: blank line.
- Line 3 onward: explanation beginning exactly with "ব্যাখ্যা:".
If the user explicitly requests no explanation, keep solution_body to the correct answer text only.

SPECIAL USER INSTRUCTIONS (CRITICAL):
{special_instructions if special_instructions else "None provided. Follow standard parsing and formatting."}

The application's response schema defines the exact fields. Return only the complete structured result; no commentary.

RAW MCQs TO PARSE:
{raw_mcq_text}
""".strip()


def parse_existing_mcqs(
    raw_mcq_text,
    user_topics,
    special_instructions,
    selected_model,
    api_key,
):
    client = genai.Client(api_key=api_key)
    allowed_topics = split_topics(user_topics)

    prompt = build_existing_mcq_prompt(
        raw_mcq_text=raw_mcq_text,
        allowed_topics=allowed_topics,
        special_instructions=special_instructions,
    )

    records = call_structured_mcq_model(
        client=client,
        model=selected_model,
        prompt=prompt,
        allowed_topics=allowed_topics,
        temperature=0.0,
        expected_counts=None,
    )

    return normalize_mcq_records(records, allowed_topics)


# ============================================================
# 8. IMAGE / HTML / DOCX HELPERS
# ============================================================

def process_html_images(text):
    """
    Convert Markdown images to embedded base64 HTML images, lock width to 2in,
    and remove leftover Pandoc dimension artifacts.
    """
    text_str = str(text)

    def replacer(match):
        img_path = match.group(1).strip()

        # Self-healing image cache.
        if "image_cache" not in st.session_state:
            st.session_state["image_cache"] = {}

        mime_type, _ = mimetypes.guess_type(img_path)
        if not mime_type:
            mime_type = "image/png"

        if img_path in st.session_state["image_cache"]:
            b64 = st.session_state["image_cache"][img_path]
            return f'<img src="data:{mime_type};base64,{b64}" width="2in" />'

        if os.path.exists(img_path):
            with open(img_path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode("utf-8")

            st.session_state["image_cache"][img_path] = b64
            return f'<img src="data:{mime_type};base64,{b64}" width="2in" />'

        # Keep unresolved paths visible instead of silently deleting them.
        return f'<img src="{img_path}" width="2in" />'

    cleaned = re.sub(
        r'!\[[^\]]*\]\(([^)]+)\)(?:\s*\{[^{}]*\})?',
        replacer,
        text_str,
        flags=re.DOTALL,
    )

    cleaned = re.sub(
        r'\{[^{}]*(?:width|height)[^{}]*\}',
        '',
        cleaned,
        flags=re.IGNORECASE | re.DOTALL,
    )

    return cleaned


def safe_newline_to_br(text):
    """Replace newlines with <br> except inside an existing HTML table."""
    text_str = str(text)

    if "<table" not in text_str.lower():
        return text_str.replace("\n", "<br>")

    parts = re.split(
        r'(<table.*?</table>)',
        text_str,
        flags=re.IGNORECASE | re.DOTALL,
    )

    for i, part in enumerate(parts):
        if not part.lower().startswith("<table"):
            parts[i] = part.replace("\n", "<br>")

    return "".join(parts)


def format_docx_layout(docx_filename):
    """
    Force A4 landscape, narrow margins, and visible borders for main and nested
    tables.
    """
    doc = Document(docx_filename)

    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

    def apply_xml_borders(tbl_obj):
        tbl = tbl_obj._tbl
        tbl_pr = tbl.tblPr

        # Force table borders.
        tbl_borders = tbl_pr.find(qn("w:tblBorders"))
        if tbl_borders is None:
            tbl_borders = OxmlElement("w:tblBorders")
            tbl_pr.append(tbl_borders)
        else:
            tbl_borders.clear()

        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "000000")
            tbl_borders.append(border)

        # Force table width to 100%.
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)

        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")

    for main_table in doc.tables:
        apply_xml_borders(main_table)

        for row in main_table.rows:
            for cell in row.cells:
                for nested_table in cell.tables:
                    apply_xml_borders(nested_table)

    doc.save(docx_filename)


def editor_data_to_records(data):
    """Handle Streamlit editor output whether it is a list or DataFrame-like."""
    if isinstance(data, list):
        return data

    if hasattr(data, "to_dict"):
        try:
            return data.to_dict("records")
        except TypeError:
            pass

    try:
        return list(data)
    except Exception:
        raise TypeError("Edited MCQ table could not be converted into records.")


def create_mcq_docx(mcq_data, output_filename="MCQs.docx"):
    records = editor_data_to_records(mcq_data)

    html = "<h1>Generated MCQs</h1>\n<table border='1'>\n"
    html += (
        "<tr>"
        "<th>Sl no.</th><th>Question Title</th><th>Option A</th>"
        "<th>Option B</th><th>Option C</th><th>Option D</th>"
        "<th>Solution Body</th><th>Correct Option</th><th>Subject</th>"
        "<th>Chapter</th><th>Topics</th><th>Question Category</th>"
        "<th>Difficulty Level</th>"
        "</tr>\n"
    )

    for mcq in records:
        html += "<tr>"
        html += f"<td>{mcq.get('sl_no', '')}</td>"

        q_title = safe_newline_to_br(mcq.get("question_title", ""))
        q_title = process_html_images(q_title)
        html += f"<td>{q_title}</td>"

        html += f"<td>{process_html_images(mcq.get('A', ''))}</td>"
        html += f"<td>{process_html_images(mcq.get('B', ''))}</td>"
        html += f"<td>{process_html_images(mcq.get('C', ''))}</td>"
        html += f"<td>{process_html_images(mcq.get('D', ''))}</td>"

        sol_body = safe_newline_to_br(mcq.get("solution_body", ""))
        sol_body = process_html_images(sol_body)
        html += f"<td>{sol_body}</td>"

        html += f"<td>{mcq.get('correct_option', '')}</td>"
        html += f"<td>{mcq.get('subject', '')}</td>"
        html += f"<td>{mcq.get('chapter', '')}</td>"
        html += f"<td>{mcq.get('topics', '')}</td>"
        html += f"<td>{mcq.get('question_category', '')}</td>"
        html += f"<td>{mcq.get('difficulty_level', '')}</td>"
        html += "</tr>\n"

    html += "</table>"

    # Pandoc converts HTML + dollar-delimited LaTeX into DOCX.
    pypandoc.convert_text(
        html,
        "docx",
        format="html+tex_math_dollars",
        outputfile=output_filename,
    )

    format_docx_layout(output_filename)

    # Do NOT delete media folders here. Keeping them allows repeated downloads
    # and preserves image access while the Streamlit session is active.


# ============================================================
# 9. STREAMLIT INTERFACE
# ============================================================

def run_mcq_interface(model_choice, api_key_input):
    # Session state
    if "user_api_key" not in st.session_state:
        st.session_state["user_api_key"] = ""
    if "mcq_data_t1" not in st.session_state:
        st.session_state["mcq_data_t1"] = None
    if "mcq_data_t2" not in st.session_state:
        st.session_state["mcq_data_t2"] = None
    if "image_cache" not in st.session_state:
        st.session_state["image_cache"] = {}

    tab1, tab2 = st.tabs(["✨ Generate New MCQs", "📋 Format Existing MCQs"])

    # ================= TAB 1: GENERATE NEW MCQS =================
    with tab1:
        with st.container(border=True):
            st.markdown("#### 📄 1. Source Material (Optional)")
            st.caption(
                "Leave blank to generate questions based purely on the topics provided below."
            )

            col_file, col_text = st.columns(2)

            with col_file:
                uploaded_file = st.file_uploader(
                    "Upload source PDF or DOCX",
                    type=["pdf", "docx"],
                    key="t1_source_file",
                )

            with col_text:
                raw_text = st.text_area(
                    "Or paste source text here",
                    height=100,
                    key="t1_source_text",
                )

        with st.container(border=True):
            st.markdown("#### 🎯 2. Question Parameters")

            topics_input = st.text_input(
                "Allowed Topics for Table Mapping (comma-separated)",
                placeholder="e.g., Hardware, Memory, Super Computers",
            )

            custom_instructions = st.text_area(
                "Custom Generation Instructions (Optional)",
                placeholder=(
                    "e.g., 'Focus heavily on numerical problems', "
                    "'Generate questions in Bengali'"
                ),
                height=100,
                key="t1_custom_instructions",
            )

        with st.container(border=True):
            st.markdown("#### 📊 3. Difficulty Breakdown")
            col1, col2, col3 = st.columns(3)

            with col1:
                num_easy = st.number_input(
                    "🟢 Easy",
                    min_value=0,
                    max_value=50,
                    value=10,
                )

            with col2:
                num_medium = st.number_input(
                    "🟡 Medium",
                    min_value=0,
                    max_value=50,
                    value=15,
                )

            with col3:
                num_hard = st.number_input(
                    "🔴 Hard",
                    min_value=0,
                    max_value=50,
                    value=5,
                )

        total_q = int(num_easy + num_medium + num_hard)

        st.markdown("<br>", unsafe_allow_html=True)
        col_btn, col_msg = st.columns([1, 2])

        with col_btn:
            generate_btn = st.button(
                "🚀 Generate MCQs",
                use_container_width=True,
                type="primary",
            )

        with col_msg:
            st.info(f"Target Generation: **{total_q}** Questions")

        if generate_btn:
            if not api_key_input:
                st.error("Please enter your API Key in the sidebar first!")
            elif total_q == 0:
                st.error("Please specify at least one question to generate.")
            elif not topics_input:
                st.error("Please provide at least one topic for table mapping.")
            else:
                with st.spinner(
                    f"Engine running... Generating {total_q} questions using {model_choice}"
                ):
                    text_to_process = ""

                    if uploaded_file:
                        if uploaded_file.name.lower().endswith(".pdf"):
                            text_to_process = extract_text_from_pdf(uploaded_file)
                        elif uploaded_file.name.lower().endswith(".docx"):
                            text_to_process = extract_text_from_docx(uploaded_file)
                    elif raw_text:
                        text_to_process = normalize_extracted_source(raw_text)

                    try:
                        st.session_state["mcq_data_t1"] = generate_mcqs(
                            text_to_process,
                            topics_input,
                            custom_instructions,
                            int(num_easy),
                            int(num_medium),
                            int(num_hard),
                            model_choice,
                            api_key_input,
                        )
                        st.success("✨ Generation Complete!")
                    except Exception as e:
                        st.error(f"An error occurred: {e}")

        if st.session_state["mcq_data_t1"]:
            with st.container(border=True):
                st.markdown("#### 👀 Preview & Edit Generated Table")
                st.caption(
                    "Double-click any cell to edit its text. You can also add or "
                    "delete rows using the tools on the right. Changes instantly "
                    "apply to your download."
                )

                edited_data_t1 = st.data_editor(
                    st.session_state["mcq_data_t1"],
                    use_container_width=True,
                    num_rows="dynamic",
                    key="editor_t1",
                )

            try:
                create_mcq_docx(edited_data_t1, "Generated_MCQs.docx")

                with open("Generated_MCQs.docx", "rb") as file:
                    st.download_button(
                        label="📥 Download Edited Word Document (.docx)",
                        data=file,
                        file_name="Generated_MCQs.docx",
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "wordprocessingml.document"
                        ),
                        use_container_width=True,
                    )
            except Exception as e:
                st.error(f"Could not build the Word document: {e}")

    # ================= TAB 2: FORMAT EXISTING MCQS =================
    with tab2:
        with st.container(border=True):
            st.markdown("#### 📝 1. Raw Input")
            st.caption("Upload or paste your unformatted questions.")

            col_file_t2, col_text_t2 = st.columns(2)

            with col_file_t2:
                uploaded_file_t2 = st.file_uploader(
                    "Upload Raw MCQs (PDF or DOCX)",
                    type=["pdf", "docx"],
                    key="t2_file",
                )

            with col_text_t2:
                raw_text_t2 = st.text_area(
                    "Or paste raw pre-written MCQs here",
                    height=150,
                    placeholder=(
                        "1. What is CPU?\n"
                        "A. Brain\nB. Memory\nC. Output\nD. Storage\nAnswer: A"
                    ),
                    key="t2_text",
                )

        with st.container(border=True):
            st.markdown("#### 🎯 2. Formatting Parameters")

            topics_input_t2 = st.text_input(
                "Target Topics (Optional)",
                placeholder="e.g., Computer Basics, Hardware",
                key="t2_topics",
            )

            special_instructions = st.text_area(
                "Special AI Instructions (Optional)",
                placeholder=(
                    "e.g., 'Set Subject to ICT', 'Fix any Bengali spelling mistakes', "
                    "'Automatically fill in missing solution explanations'"
                ),
                height=100,
                key="t2_instructions",
            )

        st.markdown("<br>", unsafe_allow_html=True)

        if st.button(
            "🛠️ Format Existing MCQs",
            use_container_width=True,
            type="primary",
            key="btn_t2",
        ):
            if not api_key_input:
                st.error("Please configure your API Key in the sidebar.")
            elif not uploaded_file_t2 and not raw_text_t2:
                st.error(
                    "Please provide your raw MCQs in the text box or upload a document."
                )
            else:
                with st.spinner(
                    f"Parsing and reformatting questions using {model_choice}..."
                ):
                    text_to_process = ""

                    if uploaded_file_t2:
                        if uploaded_file_t2.name.lower().endswith(".pdf"):
                            text_to_process = extract_text_from_pdf(uploaded_file_t2)
                        elif uploaded_file_t2.name.lower().endswith(".docx"):
                            text_to_process = extract_text_from_docx(uploaded_file_t2)
                    elif raw_text_t2:
                        text_to_process = normalize_extracted_source(raw_text_t2)

                    try:
                        st.session_state["mcq_data_t2"] = parse_existing_mcqs(
                            text_to_process,
                            topics_input_t2,
                            special_instructions,
                            model_choice,
                            api_key_input,
                        )
                        st.success("✨ Successfully reformatted into table format!")
                    except Exception as e:
                        st.error(f"Error parsing MCQs: {e}")

        if st.session_state["mcq_data_t2"]:
            with st.container(border=True):
                st.markdown("#### 👀 Preview & Edit Formatted Table")
                st.caption(
                    "Double-click any cell to edit its text. You can also add or "
                    "delete rows. Changes instantly apply to your download."
                )

                edited_data_t2 = st.data_editor(
                    st.session_state["mcq_data_t2"],
                    use_container_width=True,
                    num_rows="dynamic",
                    key="editor_t2",
                )

            try:
                create_mcq_docx(
                    edited_data_t2,
                    "Formatted_Ready_MCQs.docx",
                )

                with open("Formatted_Ready_MCQs.docx", "rb") as f:
                    st.download_button(
                        label="📥 Download Edited Word Document (.docx)",
                        data=f,
                        file_name="Formatted_Ready_MCQs.docx",
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "wordprocessingml.document"
                        ),
                        use_container_width=True,
                        key="dl_t2",
                    )
            except Exception as e:
                st.error(f"Could not build the Word document: {e}")
